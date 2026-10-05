"""Website email finder: visit company websites and collect their HR and careers addresses.

The user gives a list of companies (name optional, website required), typed in or uploaded as a CSV/Excel
sheet. A background run visits each site's pages (same site only, robots.txt respected, careers/contact/about
pages first) and keeps addresses whose name part is about hiring: hr@, careers@, career@, hr.team@,
jobs@, recruitment@, talent@ and so on. The user picks which ones to save as contacts.

Every request goes only to public internet addresses (checked again on each redirect), pages are read up to
a size cap, and each run is bounded by sites, pages per site and total time.
"""

import html
import ipaddress
import re
import socket
import threading
import time
from concurrent.futures import ThreadPoolExecutor
from urllib.parse import urljoin, urlsplit, urlunsplit
from urllib.robotparser import RobotFileParser

import requests
from flask import Blueprint, jsonify, request

import bridge
from bridge import login_required

bp = Blueprint("finder", __name__)

MAX_SITES = 100           # per run
MAX_PAGES = 25            # per site
SITES_PER_DAY = 300       # per account
PAGE_BYTES = 1_500_000
RUN_SECONDS = 30 * 60
UA = "ReachoutEmailFinder/1.0 (+https://reachout-web.netlify.app; looks for published hiring contacts)"
HINTS = ("career", "job", "hiring", "recruit", "talent", "join", "work-with", "workwithus", "vacanc", "opening",
         "contact", "about", "team", "people", "hr", "reach")
COMMON_PATHS = ("/careers", "/career", "/jobs", "/contact", "/contact-us", "/contactus", "/about", "/about-us",
                "/join-us", "/work-with-us")
SKIP_EXT = re.compile(r"\.(png|jpe?g|gif|webp|svg|ico|pdf|zip|rar|gz|mp4|mp3|webm|avi|mov|docx?|xlsx?|pptx?|css|js|"
                      r"json|xml|woff2?|ttf|eot)(\?|$)", re.I)
EMAIL_IN_TEXT = re.compile(r"[a-z0-9][a-z0-9._%+\-]{0,63}@[a-z0-9](?:[a-z0-9\-]{0,61}[a-z0-9])?(?:\.[a-z0-9\-]{1,63})*\.[a-z]{2,24}", re.I)
OBFUSCATED = re.compile(r"([a-z0-9._%+\-]{1,64})\s*[\[\(\{]\s*at\s*[\]\)\}]\s*([a-z0-9\-]{1,63}(?:\s*[\[\(\{]\s*dot\s*[\]\)\}]\s*[a-z0-9\-]{1,63})+)", re.I)
CF_EMAIL = re.compile(r'data-cfemail="([0-9a-f]{6,})"', re.I)
HREF = re.compile(r"""href\s*=\s*["']([^"'<>\s]+)["']""", re.I)
NOT_ADDRESSES = re.compile(r"\.(png|jpe?g|gif|webp|svg|css|js)$|@\dx\.|example\.(com|org)$|sentry|wixpress|domain\.com$", re.I)

RUNS = {}                 # uid -> Run (this process only; results are also saved to the workspace)
RUNS_LOCK = threading.Lock()


def C():
    return bridge.C


# ---------------------------------------------------------------- what counts as a hiring address

def hiring_kind(email):
    """'hr', 'careers' or 'hiring' when the name part of the address is about hiring; '' otherwise.

    "hr" has to stand on its own (hr@, hr.team@, hrd@, hr2@, teamhr@), so names like chris@ or
    shreya@ don't count."""
    local = email.split("@", 1)[0].lower()
    if "career" in local:
        return "careers"
    tokens = [t for t in re.split(r"[^a-z0-9]+", local) if t]
    if "humanresource" in local.replace(".", "").replace("_", "").replace("-", "") or any(
            re.fullmatch(r"hr[a-z]{0,6}\d*|[a-z]{2,6}hr\d*|hr\d*", t) for t in tokens):
        return "hr"
    if any(w in local for w in ("recruit", "talent", "hiring")) or any(t in ("jobs", "job", "joinus", "workwithus") for t in tokens):
        return "hiring"
    return ""


def extract_emails(text):
    """Every address published on a page: plain text, mailto links, [at]/(dot) spellings, Cloudflare-protected."""
    text = html.unescape(text)
    found = {m.group(0).lower().strip(".") for m in EMAIL_IN_TEXT.finditer(text)}
    for user, domain in OBFUSCATED.findall(text):
        found.add(f"{user}@{re.sub(r'\s*[\[\(\{]\s*dot\s*[\]\)\}]\s*', '.', domain)}".lower())
    for blob in CF_EMAIL.findall(text):
        try:
            key = int(blob[:2], 16)
            found.add("".join(chr(int(blob[i:i + 2], 16) ^ key) for i in range(2, len(blob), 2)).lower())
        except ValueError:
            pass
    return {e for e in found if C().EMAIL_RE.match(e) and not NOT_ADDRESSES.search(e) and len(e) <= 254}


# ---------------------------------------------------------------- fetching, only from the public internet

def host_is_public(host):
    try:
        infos = socket.getaddrinfo(host, None)
    except (OSError, UnicodeError):
        return False
    for info in infos:
        ip = ipaddress.ip_address(info[4][0].split("%")[0])
        if not ip.is_global or ip.is_multicast:
            return False
    return bool(infos)


def normalise_site(value):
    """'acme.com', 'www.acme.com/careers', 'https://acme.com' -> 'https://acme.com/...' or ''."""
    value = str(value or "").strip()
    if not value or " " in value or len(value) > 500:
        return ""
    if not re.match(r"^https?://", value, re.I):
        value = "https://" + value
    parts = urlsplit(value)
    host = (parts.hostname or "").lower().strip(".")
    if not host or "." not in host or not re.fullmatch(r"[a-z0-9.\-]+", host) or parts.port not in (None, 80, 443):
        return ""
    return urlunsplit(("https" if parts.scheme.lower() == "https" else "http", host, parts.path or "/", parts.query, ""))


def fetch(url, session, max_bytes=PAGE_BYTES):
    """GET a page; follows up to 5 redirects, each one re-checked. Returns (final_url, text) or (url, None)."""
    for _ in range(6):
        parts = urlsplit(url)
        if parts.scheme not in ("http", "https") or parts.port not in (None, 80, 443) or not host_is_public(parts.hostname or ""):
            return url, None
        try:
            r = session.get(url, timeout=(6, 12), allow_redirects=False, stream=True,
                            headers={"User-Agent": UA, "Accept": "text/html,application/xhtml+xml,text/plain;q=0.8"})
        except requests.RequestException:
            if parts.scheme == "https" and url.count("/") <= 3:  # bare https failed: some small sites are http only
                url = urlunsplit(("http",) + tuple(parts[1:]))
                continue
            return url, None
        with r:
            if r.is_redirect and r.headers.get("location"):
                url = urljoin(url, r.headers["location"])
                continue
            kind = r.headers.get("content-type", "")
            if r.status_code >= 400 or not any(k in kind for k in ("html", "text/plain", "xml")) and kind:
                return url, None
            data = b""
            for chunk in r.iter_content(65536):
                data += chunk
                if len(data) >= max_bytes:
                    break
            return url, data.decode(r.encoding or "utf-8", errors="replace")
    return url, None


def same_site(host, base):
    host, base = host.lower().removeprefix("www."), base.lower().removeprefix("www.")
    return host == base or host.endswith("." + base)


def scan_site(site, stop, pages_cap=MAX_PAGES):
    """Visit one website; returns {'url', 'pages', 'emails': [{email, kind, page}], 'other': n, 'error'}."""
    out = {"url": site, "pages": 0, "emails": [], "other": 0, "error": ""}
    session = requests.Session()
    final, page = fetch(site, session)
    if page is None:
        out["error"] = "Couldn't open this website (it may be down or block automated visits)."
        return out
    base = urlsplit(final).hostname or ""
    root = f"{urlsplit(final).scheme}://{urlsplit(final).netloc}"
    robots = RobotFileParser()
    _, rtxt = fetch(root + "/robots.txt", session, 200_000)
    robots.parse((rtxt or "").splitlines())
    seen, queue, found = {final}, [], {}
    queue += [root + p for p in COMMON_PATHS]
    all_emails = set()

    def take(url, text):
        out["pages"] += 1
        for e in extract_emails(text):
            all_emails.add(e)
            kind = hiring_kind(e)
            if kind and e not in found:
                found[e] = {"email": e, "kind": kind, "page": url, "same_domain": same_site(e.split("@")[1], base)}
        links = []
        for href in HREF.findall(text):
            if href.lower().startswith("mailto:"):
                continue
            u = urljoin(url, html.unescape(href)).split("#")[0]
            p = urlsplit(u)
            if p.scheme in ("http", "https") and same_site(p.hostname or "", base) and not SKIP_EXT.search(p.path) and u not in seen:
                links.append(u)
        # Pages that usually list hiring contacts go first.
        links.sort(key=lambda u: -sum(h in u.lower() for h in HINTS))
        queue[:0] = [u for u in links if any(h in u.lower() for h in HINTS)]
        queue.extend(u for u in links if not any(h in u.lower() for h in HINTS))

    take(final, page)
    while queue and out["pages"] < pages_cap and not stop.is_set():
        url = queue.pop(0)
        if url in seen:
            continue
        seen.add(url)
        if not robots.can_fetch(UA, url) or not robots.can_fetch("*", url):
            continue
        got, text = fetch(url, session)
        if text is not None and same_site(urlsplit(got).hostname or "", base):
            seen.add(got)
            take(got, text)
        time.sleep(0.4)  # be gentle with small sites
    out["emails"] = sorted(found.values(), key=lambda x: (not x["same_domain"], {"hr": 0, "careers": 1, "hiring": 2}[x["kind"]], x["email"]))
    out["other"] = len(all_emails) - len(found)
    return out


# ---------------------------------------------------------------- runs

class Run:
    def __init__(self, uid, items):
        self.uid, self.items = uid, items
        self.stop = threading.Event()
        self.results = []
        self.done = 0
        self.state = "running"
        self.started = time.time()

    def snapshot(self):
        return {"state": self.state, "total": len(self.items), "done": self.done, "started": self.started,
                "results": self.results}


def run_scan(run):
    ws = C().Workspace(run.uid)

    def one(item):
        if run.stop.is_set() or time.time() - run.started > RUN_SECONDS:
            return {**item, "pages": 0, "emails": [], "other": 0, "error": "Skipped (stopped or time limit reached)."}
        try:
            return {**item, **scan_site(item["website"], run.stop)}
        except Exception as e:  # one broken site never ends the whole run
            print(f"[Reachout] finder: {item['website']}: {e}", flush=True)
            return {**item, "pages": 0, "emails": [], "other": 0, "error": "Couldn't read this website."}

    try:
        with ThreadPoolExecutor(max_workers=4) as pool:
            for res in pool.map(one, run.items):
                res.pop("url", None)
                run.results.append(res)
                run.done += 1
        run.state = "stopped" if run.stop.is_set() else "done"
    except Exception as e:
        print(f"[Reachout] finder run failed: {e}", flush=True)
        run.state = "error"
    with ws.lock:
        ws.save("finder", run.snapshot())
    found = sum(len(r["emails"]) for r in run.results)
    from features.notify import notify
    notify(run.uid, f"Website finder: {found} hiring address{'es' if found != 1 else ''} found",
                  f"Checked {run.done} website{'s' if run.done != 1 else ''}. Pick the ones to save as contacts.",
                  "/app/finder", "info")


SITE_HEADERS = {"website", "websites", "url", "urls", "site", "domain", "web", "link", "homepage", "company website", "career page", "careers page"}
NAME_HEADERS = {"company", "company name", "name", "organisation", "organization", "employer"}


def read_grid(core, filename, data):
    """Rows of cells from a .csv or .xlsx file (first 1000 rows)."""
    import csv
    import io
    import zipfile
    name = filename.lower()
    try:
        if name.endswith((".xlsx", ".xlsm")):
            from openpyxl import load_workbook
            with zipfile.ZipFile(io.BytesIO(data)) as z:  # check what it expands to before opening it
                if len(z.infolist()) > 200 or sum(i.file_size for i in z.infolist()) > 40 * 1024 * 1024:
                    raise core.Invalid("That spreadsheet is too large. Keep just the sheet with the websites.", "file")
            sheet = load_workbook(io.BytesIO(data), read_only=True, data_only=True).active
            grid = [list(r[:10]) for _, r in zip(range(1000), sheet.iter_rows(values_only=True))]
        elif name.endswith((".csv", ".txt")):
            text = data.decode("utf-8-sig", errors="replace")
            grid = [r[:10] for _, r in zip(range(1000), csv.reader(io.StringIO(text)))]
        else:
            raise core.Invalid("Upload a CSV or Excel (.xlsx) file.", "file")
    except core.Invalid:
        raise
    except Exception:
        raise core.Invalid("We couldn't read that file. Check it opens in Excel, then try again.", "file")
    return [r for r in grid if any(str(c or "").strip() for c in r)]


def read_items(core):
    """Companies from the JSON body or an uploaded CSV/Excel sheet: [{'company', 'website'}]."""
    if request.files.get("file"):
        f = request.files["file"]
        data = f.read(2 * 1024 * 1024 + 1)
        if len(data) > 2 * 1024 * 1024:
            raise core.Invalid("That file is larger than 2 MB.", "file")
        grid = read_grid(core, f.filename or "", data)
        raw = []
        head = [str(h or "").strip().lower() for h in grid[0]] if grid else []
        site_col = next((i for i, h in enumerate(head) if h in SITE_HEADERS), None)
        name_col = next((i for i, h in enumerate(head) if h in NAME_HEADERS), None)
        body_rows = grid[1:] if site_col is not None or name_col is not None else grid  # no header: a plain list of sites
        for r in body_rows:
            cells = [str(c or "").strip() for c in r]
            if site_col is not None and site_col < len(cells):
                site = cells[site_col]
            else:  # take the first value that looks like a web address
                site = next((c for c in cells if re.search(r"[a-z0-9-]+\.[a-z]{2,}", c, re.I) and "@" not in c), "")
            company = cells[name_col] if name_col is not None and name_col < len(cells) else ""
            raw.append({"company": company, "website": site})
    else:
        raw = (core.body().get("items") or [])[:MAX_SITES * 3]
    items, seen, bad = [], set(), 0
    for r in raw:
        if not isinstance(r, dict):
            continue
        site = normalise_site(r.get("website"))
        if not site:
            bad += 1 if str(r.get("website") or "").strip() else 0
            continue
        host = urlsplit(site).hostname.removeprefix("www.")
        if host in seen:
            continue
        seen.add(host)
        items.append({"company": str(r.get("company") or "").strip()[:120] or host.split(".")[0].capitalize(), "website": site})
    if not items:
        raise core.Invalid("Add at least one company website, like acme.com." if not bad
                           else "None of those website addresses look right. Use addresses like acme.com.", "website")
    if len(items) > MAX_SITES:
        raise core.Invalid(f"One run can check up to {MAX_SITES} websites. Split the list and run it in parts.", "file")
    return items, bad


@bp.post("/api/finder/scan")
@login_required
def start_scan(ws):
    core = C()
    items, bad = read_items(core)
    with RUNS_LOCK:
        cur = RUNS.get(ws.uid)
        if cur and cur.state == "running":
            raise core.Invalid("A search is already running. Wait for it to finish or stop it.", status=409)
        for _ in items:
            if core.rate_limited(("finder", ws.uid), SITES_PER_DAY, 86400):
                raise core.Invalid(f"You can check up to {SITES_PER_DAY} websites a day. Try again tomorrow.", status=429)
        run = RUNS[ws.uid] = Run(ws.uid, items)
    threading.Thread(target=run_scan, args=(run,), daemon=True, name="finder").start()
    return jsonify(ok=True, total=len(items), skipped=bad)


@bp.get("/api/finder")
@login_required
def status(ws):
    run = RUNS.get(ws.uid)
    snap = run.snapshot() if run else (ws.load("finder", None) or {"state": "idle", "total": 0, "done": 0, "results": []})
    emails = {r.get("email", "").lower() for r in ws.load("recipients", [])}
    for site in snap.get("results", []):
        for e in site.get("emails", []):
            e["saved"] = e["email"] in emails
    return jsonify(snap | {"max_sites": MAX_SITES})


@bp.post("/api/finder/stop")
@login_required
def stop(ws):
    run = RUNS.get(ws.uid)
    if run and run.state == "running":
        run.stop.set()
    return jsonify(ok=True)


@bp.post("/api/finder/save")
@login_required
def save_contacts(ws):
    """Save picked addresses as contacts in a list (default "Website finder"); existing emails are skipped."""
    core = C()
    p = core.body()
    picks = p.get("items") or []
    list_name = core.v_text(p.get("list") or "Website finder", "list", "List", 60)
    if not isinstance(picks, list) or not picks:
        raise core.Invalid("Pick at least one address to save.")
    added = duplicates = invalid = 0
    with ws.lock:
        rows = ws.load("recipients", [])
        have = {r.get("email", "").lower() for r in rows}
        for it in picks[:1000]:
            if not isinstance(it, dict):
                continue
            try:
                email = core.v_email(it.get("email"))
                company = core.v_text(it.get("company"), "company", "Company", 120)
                site = normalise_site(it.get("website"))
            except core.Invalid:
                invalid += 1
                continue
            if email in have:
                duplicates += 1
                continue
            if len(rows) >= core.MAX_CONTACTS:
                raise core.Invalid(f"You've reached the limit of {core.MAX_CONTACTS} contacts.")
            kind = hiring_kind(email)
            rows.append({"id": core.new_id(), "added_at": core.now_iso(), "name": f"{company} {'HR' if kind == 'hr' else 'Careers'}".strip(),
                         "company": company, "email": email, "list": list_name, "stage": "new",
                         "email_type": "careers" if kind else "general", "website": site})
            have.add(email)
            added += 1
        ws.save("recipients", rows)
    return jsonify(added=added, duplicates=duplicates, invalid=invalid)


@bp.get("/api/finder/sample")
@login_required
def sample(ws):
    from flask import Response
    return Response("company,website\nAcme Corp,acme.com\nNorthwind,https://www.northwind.example\n", mimetype="text/csv",
                    headers={"Content-Disposition": "attachment; filename=websites-sample.csv"})
